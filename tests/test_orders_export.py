"""The export of orders to an Excel file: what the page lists, nothing written anywhere, and nothing in a cell that Excel could run (owner, 2026-10-15/19)."""
from __future__ import annotations

import io
import re
from datetime import datetime

import openpyxl

from panel_helpers import login, make_app
from test_orders_core import NOW, TERMS, item
from tires_sync import order_points as points, orders, orders_export, store

HEAD = ["№ заказа", "Создан", "Этап", "Клиент", "Телефон", "Точка", "Ответственный", "Источник", "Теги", "Позиции", "Штук", "Сумма, ₽",
        "Дата доставки", "Номер заказа WebMIM", "Статус WebMIM", "Комментарий", "Чат Авито"]


def db(tmp_path):
    conn = store.connect(str(tmp_path / "x.db"))
    orders.ensure_pipeline(conn, NOW)
    return conn


def sheet(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data)).active


def test_the_file_has_a_header_one_row_per_order_newest_first_and_the_positions_and_the_sum(tmp_path):
    conn = db(tmp_path)
    pid = points.save_point(conn, {"name": "Гамма"})
    first = orders.create_card(conn, {"client_name": "Иван", "client_phone": "+79000000017", "point_id": pid, "tags": "срочно,Основной", "comment": "позвонить"}, NOW, "worker2")
    item(conn, first, name="Pirelli 225/55 R17", qty=4, sale_unit=9600.0)
    item(conn, first, goods_id="G2", name="Ikon 205/55 R16", qty=2, sale_unit=5000.0)
    second = orders.create_card(conn, {"client_name": "Пётр"}, NOW, "worker2")
    rows = list(sheet(orders_export.build(conn, datetime(2026, 10, 15, 12, 0))).iter_rows(values_only=True))
    assert list(rows[0]) == HEAD and [r[0] for r in rows[1:]] == [second, first]                      # the newest first
    one = dict(zip(HEAD, rows[2]))
    assert one["Клиент"] == "Иван" and one["Телефон"] == "+79000000017" and one["Точка"] == "Гамма" and one["Этап"] == "Новый заказ"
    assert one["Позиции"] == "Pirelli 225/55 R17 × 4; Ikon 205/55 R16 × 2" and one["Штук"] == 6 and one["Сумма, ₽"] == 4 * 9600 + 2 * 5000
    assert one["Теги"] == "срочно, Основной" and one["Комментарий"] == "позвонить" and one["Ответственный"] == "worker2"
    assert dict(zip(HEAD, rows[1]))["Позиции"] in ("", None)


def test_the_filters_of_the_page_cut_the_file_the_same_way(tmp_path):
    conn = db(tmp_path)
    pid = points.save_point(conn, {"name": "Гамма"})
    a = orders.create_card(conn, {"client_name": "Иван", "point_id": pid, "tags": "срочно"}, NOW, "o")
    b = orders.create_card(conn, {"client_name": "Пётр"}, NOW, "o")
    now = datetime(2026, 10, 15, 12, 0)
    ids = lambda **kw: [r[0] for r in list(sheet(orders_export.build(conn, now, **kw)).iter_rows(values_only=True))[1:]]
    assert ids() == [b, a] and ids(point_id=pid) == [a] and ids(tag="срочно") == [a] and ids(query="Пётр") == [b] and ids(query="нет такого") == []


def test_nothing_in_a_cell_can_be_run_by_excel_and_forbidden_signs_do_not_break_the_file(tmp_path):
    conn = db(tmp_path)
    evil = '=HYPERLINK("http://evil.example","x")'
    cid = orders.create_card(conn, {"client_name": evil, "comment": "+SUM(1,1)\x00\x07 дальше", "tags": "@x"}, NOW, "o")
    item(conn, cid, name="=cmd|' /C calc'!A0")
    ws = sheet(orders_export.build(conn, datetime(2026, 10, 15, 12, 0)))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            assert cell.data_type != "f"                                                                    # no formula anywhere
    by = {h: c for h, c in zip(HEAD, ws[2])}
    assert by["Клиент"].value == evil and by["Клиент"].data_type == "s" and by["Комментарий"].value == "+SUM(1,1) дальше"
    assert by["Позиции"].data_type == "s" and by["Позиции"].value.startswith("=cmd")


def test_the_page_gives_the_file_to_a_logged_in_manager_with_the_filters_and_a_button(tmp_path):
    app = make_app(tmp_path, NOW=lambda: datetime(2026, 10, 15, 9, 0, 0))
    client = app.test_client()
    assert app.test_client().get("/orders/export.xlsx").status_code == 302
    login(client)
    conn = store.connect(app.config["DB_PATH"])
    orders.ensure_pipeline(conn, NOW)
    orders.create_card(conn, {"client_name": "Иван"}, NOW, "o")
    orders.create_card(conn, {"client_name": "Пётр"}, NOW, "o")
    resp = client.get("/orders/export.xlsx?q=Иван")
    assert resp.status_code == 200 and resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "orders-2026-10-15.xlsx" in resp.headers["Content-Disposition"] and resp.headers["Cache-Control"].startswith("no-store")
    assert [r[3] for r in list(sheet(resp.get_data()).iter_rows(values_only=True))[1:]] == ["Иван"]
    page = client.get("/orders?q=Иван").get_data(as_text=True)
    assert re.search(r'href="/orders/export\.xlsx\?p=&amp;t=&amp;q=%D0%98%D0%B2%D0%B0%D0%BD"', page) and "Выгрузить в Excel" in page     # the button carries the filters of the page


def test_a_machine_without_the_excel_library_gets_a_message_not_an_error_page(tmp_path, monkeypatch):
    app = make_app(tmp_path, NOW=lambda: datetime(2026, 10, 15, 9, 0, 0))
    client = app.test_client()
    login(client)

    def broken(*a, **k):
        raise ImportError("no openpyxl")
    monkeypatch.setattr(orders_export, "build", broken)
    resp = client.get("/orders/export.xlsx")
    assert resp.status_code == 302 and resp.headers["Location"].endswith("/orders")
    assert "не установлена библиотека openpyxl" in client.get("/orders").get_data(as_text=True)


def test_the_two_signs_xml_forbids_do_not_break_the_file(tmp_path):
    conn = db(tmp_path)
    orders.create_card(conn, {"client_name": "Иван￾￿", "comment": "a￾b"}, NOW, "o")
    ws = sheet(orders_export.build(conn, datetime(2026, 10, 15, 12, 0)))
    by = {h: c.value for h, c in zip(HEAD, ws[2])}
    assert by["Клиент"] == "Иван" and by["Комментарий"] == "ab"


def test_the_excel_library_is_imported_only_when_a_file_is_asked_for(tmp_path, monkeypatch):
    import sys
    app = make_app(tmp_path, NOW=lambda: datetime(2026, 10, 15, 9, 0, 0))
    client = app.test_client()
    login(client)
    monkeypatch.setitem(sys.modules, "openpyxl", None)                              # as on a machine where it is not installed
    resp = client.get("/orders/export.xlsx")
    assert resp.status_code == 302 and "не установлена библиотека openpyxl" in client.get("/orders").get_data(as_text=True)
    assert client.get("/orders").status_code == 200                                 # the rest of the panel does not need it


def test_nothing_about_the_purchase_or_the_margin_is_in_the_file(tmp_path):
    conn = db(tmp_path)
    cid = orders.create_card(conn, {"client_name": "Иван"}, NOW, "o")
    item(conn, cid, qty=4, sale_unit=9600.0, supplier_price=1234.56)
    data = orders_export.build(conn, datetime(2026, 10, 15, 12, 0))
    ws = sheet(data)
    cells = [c.value for row in ws.iter_rows() for c in row]
    assert not any(isinstance(v, str) and ("закуп" in v.lower() or "марж" in v.lower()) for v in cells)
    assert 1234.56 not in cells and 4938.24 not in cells                            # neither the price of the supplier nor four of them
