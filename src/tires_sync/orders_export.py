"""The orders as an Excel file: the same orders the page lists under the same filters, one row each, newest first. Read only: nothing is written
anywhere. Every text goes into the file as plain text (never as a formula) and without the signs Excel does not accept."""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import List, Optional

from . import orders, oth_sale

HEAD = ["№ заказа", "Создан", "Этап", "Клиент", "Телефон", "Точка", "Ответственный", "Источник", "Теги", "Позиции", "Штук", "Сумма, ₽",
        "Дата доставки", "Номер заказа WebMIM", "Статус WebMIM", "Комментарий", "Чат Авито"]
WIDTHS = [10, 17, 22, 28, 18, 16, 16, 18, 24, 60, 8, 14, 14, 18, 18, 40, 16]
MONEY = {"Сумма, ₽"}                                                # the columns of money: never in the file of a point (the owner 2026-10-27)
FORBIDDEN = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")        # signs a sheet cannot hold (XML 1.0 forbids these too)


def _text(value) -> str:
    return FORBIDDEN.sub("", str(value or ""))


def _created(value: str) -> str:
    try:
        return datetime.fromisoformat(str(value)).strftime("%d.%m.%Y %H:%M")
    except ValueError:
        return _text(value)


def _row(card: dict, hidden=frozenset()) -> list:
    items = card["items"]
    money = orders.totals(items)
    return [card["id"], _created(card["created_at"]), card["stage"]["name"], card["client_name"], card["client_phone"],
            card["point"]["name"] if card["point"] else "", card["responsible"], orders.SOURCES.get(card["source"], card["source"]),
            ", ".join(t for t in card["tags"] if t.lower() not in hidden), "; ".join("%s × %d" % (i["name"], i["qty"]) for i in items), money["qty"], money["sale"],
            card["delivery_date"] or (oth_sale.IN_STOCK if card.get("in_stock") else ""), card["webmim_number"], card["webmim_status"], card["comment"], card["avito_chat"]]


def build(conn, now: datetime, point_id: Optional[int] = None, tag: str = "", query: str = "", money: bool = True) -> bytes:
    """The file (bytes) with the orders the page would list under these filters; without `money` (a point) no column of MONEY."""
    import openpyxl
    from openpyxl.styles import Font
    data = orders.board(conn, now, point_id=point_id, tag=tag, query=query)
    ids = sorted((c["id"] for col in data["columns"] for c in col["cards"]), reverse=True)
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Заказы"
    hidden = frozenset() if money else frozenset(orders.account_tags(conn))
    # the file of a point: no tag with an account's name (the owner 2026-10-29)
    keep = [i for i, name in enumerate(HEAD) if money or name not in MONEY]
    sheet.append([HEAD[i] for i in keep])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for card_id in ids:
        card = orders.get_card(conn, card_id)
        if not money:                                                    # a point: the comment without an account's name either
            card = dict(card, comment=orders.without_account_names(conn, card["comment"]))
        values = [_row(card, hidden)[i] for i in keep]
        sheet.append([None] * len(values))
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=sheet.max_row, column=column)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.value = value
            else:
                cell.value = _text(value)
                cell.data_type = "s"                                    # a text, never a formula, whatever it starts with
    for index, width in enumerate([WIDTHS[i] for i in keep], start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
